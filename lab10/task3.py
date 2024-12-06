from docx import Document

document = Document('lab10.docx')

table = document.tables[0]

for i, type in enumerate(['Flash', 'SRAM', 'EEPROM']):
    print(f'{type:<8} {table.cell(i+1, 2).text[1:]}')

document.save('lab10.docx')